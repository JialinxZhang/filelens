# -*- coding: utf-8 -*-
"""文件内容抽取与摘要生成。

按扩展名分派到不同抽取器，产出统一结构：
  kind / kind_label  类型
  title              推断标题（第一行有意义的文字，兜底为文件名）
  gist               开头段落摘要（约 500 字）
  points             疑似小标题的行（提取式要点，非 AI 生成）
  stats              类型相关统计（页数 / 工作表 / 幻灯片 / 字符数…）
  content            正文预览文本
  warnings           解析告警（扫描件、损坏、不支持的格式等）
"""
import html
import os
import re
import struct
import zipfile

PREVIEW_CHARS = 6_000      # 默认预览长度
FULL_CHARS = 60_000        # “展开全文”长度
GIST_CHARS = 480
READ_RAW_CAP = 8 * 1024 * 1024   # 纯文本文件最多读入的字节数

TEXT_EXTS = {
    'txt', 'md', 'markdown', 'log', 'json', 'xml', 'html', 'htm', 'srt',
    'ini', 'cfg', 'conf', 'yaml', 'yml', 'toml', 'csv', 'tsv', 'js', 'ts',
    'py', 'c', 'cpp', 'h', 'java', 'css', 'bat', 'sh', 'sql', 'r', 'm',
}
IMAGE_EXTS = {'jpg', 'jpeg', 'png', 'gif', 'webp', 'bmp', 'svg', 'ico', 'tif', 'tiff', 'heic'}
ARCHIVE_EXTS = {'rar', '7z', 'gz', 'tar', 'bz2'}
MEDIA_EXTS = {'mp3', 'wav', 'flac', 'aac', 'mp4', 'avi', 'mkv', 'mov', 'wmv', 'flv', 'm4a'}

KIND_LABELS = {
    'pdf': 'PDF 文档', 'word': 'Word 文档', 'word-legacy': 'Word 97-2003',
    'excel': 'Excel 表格', 'ppt': 'PPT 演示', 'text': '文本文件',
    'image': '图片', 'zip': 'ZIP 压缩包', 'archive': '压缩包',
    'media': '音视频', 'binary': '二进制文件',
}


def kind_of(ext):
    if ext == 'pdf':
        return 'pdf'
    if ext == 'docx':
        return 'word'
    if ext == 'doc':
        return 'word-legacy'
    if ext in ('xlsx', 'xlsm', 'xls', 'et'):
        return 'excel'
    if ext == 'pptx':
        return 'ppt'
    if ext in TEXT_EXTS:
        return 'text'
    if ext in IMAGE_EXTS:
        return 'image'
    if ext == 'zip':
        return 'zip'
    if ext in ARCHIVE_EXTS:
        return 'archive'
    if ext in MEDIA_EXTS:
        return 'media'
    return 'binary'


# ---------------------------------------------------------------- 摘要启发式

_HEADING_RES = [
    re.compile(r'^#{1,6}\s+\S'),
    re.compile(r'^[一二三四五六七八九十百]+\s*[、.．:：]'),
    re.compile(r'^[（(][一二三四五六七八九十]+[)）]'),
    re.compile(r'^\d+(?:\.\d+){0,3}\s*[、.．:：]\S'),
    re.compile(r'^第[一二三四五六七八九十\d]+\s*[章节部分条]'),
    re.compile(r'^[（(]\d+[)）]\s*\S'),
]


def _looks_like_heading(line):
    if not (6 <= len(line) <= 60):
        return False
    if line.rstrip().endswith(('。', '；', ';', '，', ',')):
        return False
    return any(r.match(line) for r in _HEADING_RES)


def _clean_line(line):
    return re.sub(r'\s+', ' ', line.replace('\u3000', ' ')).strip()


def summarize_text(text, max_points=8):
    """从纯文本做提取式摘要：标题 + 开头段落 + 小标题要点。"""
    lines = [_clean_line(l) for l in text.splitlines()]
    nonempty = [l for l in lines if l]

    title = ''
    for l in nonempty[:5]:
        if len(l) <= 60 and re.search(r'[\w\u4e00-\u9fff]', l) and not re.fullmatch(r'[\d./\-\s:年月日]+', l):
            title = l
            break

    gist_parts, gist_len = [], 0
    for l in nonempty:
        if gist_len >= GIST_CHARS:
            break
        if l == title and not gist_parts:
            continue
        gist_parts.append(l)
        gist_len += len(l) + 1
    gist = '\n'.join(gist_parts[:6])

    points, seen = [], set()
    for l in nonempty[1:]:
        if len(points) >= max_points:
            break
        key = l[:24]
        if key in seen or l == title:
            continue
        if _looks_like_heading(l):
            seen.add(key)
            points.append(l.lstrip('# '))
    return title, gist, points


# ---------------------------------------------------------------- 各类型抽取

def _extract_pdf(path, limit):
    out = {'warnings': []}
    try:
        try:
            import pymupdf as pm
        except ImportError:
            import fitz as pm
        with pm.open(path) as doc:
            pages = doc.page_count
            meta = doc.metadata or {}
            n_read = min(pages, 40)
            chunks = []
            total_chars = 0
            for i in range(n_read):
                t = doc[i].get_text('text').strip()
                if t:
                    chunks.append(t)
                    total_chars += len(t)
            text = '\n'.join(chunks)
            out['stats'] = {'pages': pages, 'chars': total_chars}
            meta_title = (meta.get('title') or '').strip()
            if total_chars < 30 and pages > 0:
                out['warnings'].append('这是扫描件（图片型 PDF），未检测到文本层；已在预览中显示首页截图。')
                out['content'] = ''
                out['title'] = meta_title or None
                out['gist'], out['points'] = '', []
                return out
            truncated = len(text) > limit or n_read < pages
            out['content'] = text[:limit]
            out['content_truncated'] = truncated
            title, gist, points = summarize_text(text)
            if meta_title and len(meta_title) <= 80:
                title = meta_title
            out.update(title=title, gist=gist, points=points)
            if n_read < pages:
                out['warnings'].append(f'仅读取了前 {n_read} 页（共 {pages} 页）。')
            return out
    except ImportError:
        out['warnings'].append('未安装 PyMuPDF（pip install pymupdf），无法解析 PDF。')
    except Exception as e:
        out['warnings'].append(f'PDF 解析失败：{e}')
    out.setdefault('stats', {})
    out.setdefault('content', '')
    return out


def _unescape(s):
    return html.unescape(s)


def _extract_docx(path, limit):
    out = {'warnings': []}
    try:
        with zipfile.ZipFile(path) as z:
            xml = z.read('word/document.xml').decode('utf-8', 'replace')
            paras = []
            for p_xml in xml.split('</w:p>'):
                # <w:t> 或 <w:t xml:space="preserve">，不能误匹配 <w:tbl>/<w:tc> 等
                ts = re.findall(r'<w:t(?:\s[^>]*)?>(.*?)</w:t>', p_xml, re.S)
                if ts:
                    line = _unescape(''.join(ts)).strip()
                    if line:
                        paras.append(line)
            text = '\n'.join(paras)
            title, gist, points = summarize_text(text)
            try:
                core = z.read('docProps/core.xml').decode('utf-8', 'replace')
                m = re.search(r'<dc:title>(.*?)</dc:title>', core, re.S)
                if m and _unescape(m.group(1)).strip():
                    title = _unescape(m.group(1)).strip()
            except KeyError:
                pass
            out.update(content=text[:limit], content_truncated=len(text) > limit,
                       title=title, gist=gist, points=points,
                       stats={'paragraphs': len(paras), 'chars': len(text)})
    except Exception as e:
        out['warnings'].append(f'docx 解析失败：{e}')
        out.setdefault('stats', {})
        out.setdefault('content', '')
    return out


def _fmt_cell(v):
    if v is None:
        return ''
    if isinstance(v, float):
        return str(round(v, 6)).rstrip('0').rstrip('.')
    return str(v)


def _extract_xlsx(path, limit):
    out = {'warnings': []}
    try:
        import openpyxl
        wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
        names = wb.sheetnames
        blocks, total_cells = [], 0
        for name in names[:4]:
            ws = wb[name]
            rows = []
            for i, row in enumerate(ws.iter_rows(values_only=True)):
                cells = [_fmt_cell(v)[:40] for v in row[:10]]
                if any(cells):
                    rows.append(' | '.join(cells))
                total_cells += 1
                if i >= 14 or total_cells > 200:
                    break
            blocks.append(f'【{name}】\n' + ('\n'.join(rows) or '（空工作表）'))
        wb.close()
        content = '\n\n'.join(blocks)
        out.update(content=content[:limit], content_truncated=len(content) > limit,
                   title=names[0] if names else None, gist=content[:GIST_CHARS],
                   points=[f'共 {len(names)} 个工作表：' + '、'.join(names[:6]) + ('…' if len(names) > 6 else '')],
                   stats={'sheets': len(names), 'sheet_names': names[:12]})
    except Exception as e:
        out['warnings'].append(f'xlsx 解析失败：{e}')
        out.setdefault('stats', {})
        out.setdefault('content', '')
    return out


def _extract_pptx(path, limit):
    out = {'warnings': []}
    try:
        with zipfile.ZipFile(path) as z:
            slide_names = sorted(
                (n for n in z.namelist() if re.fullmatch(r'ppt/slides/slide\d+\.xml', n)),
                key=lambda n: int(re.search(r'(\d+)', n).group(1)))
            blocks = []
            for idx, n in enumerate(slide_names[:8], 1):
                xml = z.read(n).decode('utf-8', 'replace')
                ts = [_unescape(t).strip() for t in re.findall(r'<a:t(?:\s[^>]*)?>(.*?)</a:t>', xml, re.S)]
                ts = [t for t in ts if t]
                if ts:
                    blocks.append(f'【幻灯片 {idx}】\n' + '\n'.join(ts))
            content = '\n'.join(blocks)
            first_lines = [l for b in blocks for l in b.splitlines()[1:3]]
            title = first_lines[0] if first_lines else None
            out.update(content=content[:limit], content_truncated=len(slide_names) > 8,
                       title=title, gist='\n'.join(first_lines[:8]),
                       points=[l for l in first_lines[1:9]],
                       stats={'slides': len(slide_names)})
    except Exception as e:
        out['warnings'].append(f'pptx 解析失败：{e}')
        out.setdefault('stats', {})
        out.setdefault('content', '')
    return out


def _extract_zip(path, limit):
    out = {'warnings': []}
    try:
        with zipfile.ZipFile(path) as z:
            infos = z.infolist()
            lines = []
            for info in infos[:300]:
                name = info.filename
                if not info.flag_bits & 0x800:   # 无 UTF-8 标志位：macOS 压缩包多为 UTF-8 字节，国产压缩软件多为 GBK
                    raw = None
                    try:
                        raw = info.filename.encode('cp437')
                    except UnicodeEncodeError:
                        pass
                    if raw:
                        for enc in ('utf-8', 'gb18030'):
                            try:
                                name = raw.decode(enc)
                                break
                            except UnicodeDecodeError:
                                continue
                size = f'{info.file_size / 1024 / 1024:.1f} MB' if info.file_size > 1024 * 1024 else f'{max(info.file_size, 0) / 1024:.0f} KB'
                lines.append(f'{name}　({size})')
            content = '\n'.join(lines) or '（空压缩包）'
            total = sum(i.file_size for i in infos)
            out.update(content=content[:limit], content_truncated=len(infos) > 300,
                       title=None, gist='\n'.join(lines[:12]), points=[],
                       stats={'entries': len(infos), 'uncompressed': total},
                       warnings=(['条目过多，仅列出前 300 项。'] if len(infos) > 300 else []))
    except Exception as e:
        out['warnings'].append(f'zip 解析失败：{e}')
        out.setdefault('stats', {})
        out.setdefault('content', '')
    return out


def _decode_text(raw):
    for enc in ('utf-8', 'gb18030'):
        try:
            return raw.decode(enc)
        except UnicodeDecodeError:
            continue
    return raw.decode('utf-8', 'replace')


def _extract_text(path, limit):
    out = {'warnings': []}
    try:
        with open(path, 'rb') as f:
            raw = f.read(READ_RAW_CAP)
        text = _decode_text(raw).lstrip('\ufeff')
        if path.lower().endswith(('.html', '.htm')):
            text = re.sub(r'(?is)<(script|style)[^>]*>.*?</\1>', '', text)
            text = re.sub(r'(?s)<[^>]+>', ' ', text)
            text = html.unescape(re.sub(r'[ \t]+', ' ', text))
        title, gist, points = summarize_text(text)
        out.update(content=text[:limit], content_truncated=len(text) > limit,
                   title=title, gist=gist, points=points,
                   stats={'chars': len(text)},
                   warnings=(['文件较大，仅读取前 8 MB。'] if len(raw) >= READ_RAW_CAP else []))
    except Exception as e:
        out['warnings'].append(f'读取失败：{e}')
        out.setdefault('stats', {})
        out.setdefault('content', '')
    return out


def _image_size(path):
    """轻量读取 PNG/JPEG/GIF 尺寸，失败返回 None。"""
    try:
        with open(path, 'rb') as f:
            head = f.read(32)
        if head[:8] == b'\x89PNG\r\n\x1a\n':
            w, h = struct.unpack('>II', head[16:24])
            return w, h
        if head[:6] in (b'GIF87a', b'GIF89a'):
            w, h = struct.unpack('<HH', head[6:10])
            return w, h
        if head[:2] == b'\xff\xd8':
            with open(path, 'rb') as f:
                data = f.read(256 * 1024)
            for m in re.finditer(rb'\xff[\xc0-\xcf]', data):
                seg = data[m.start():m.start() + 9]
                if len(seg) >= 9 and seg[1] not in (0xc4, 0xc8, 0xcc):
                    h, w = struct.unpack('>HH', seg[5:9])
                    return w, h
    except Exception:
        pass
    return None


# ---------------------------------------------------------------- 入口

def inspect(path, full=False):
    ext = os.path.splitext(path)[1][1:].lower()
    kind = kind_of(ext)
    limit = FULL_CHARS if full else PREVIEW_CHARS
    base = {
        'kind': kind,
        'kind_label': KIND_LABELS.get(kind, ext.upper() or '未知'),
        'title': None, 'gist': '', 'points': [], 'stats': {},
        'content': '', 'content_truncated': False, 'warnings': [],
        'image_size': None,
    }

    if kind == 'pdf':
        base.update(_extract_pdf(path, limit))
    elif kind == 'word':
        base.update(_extract_docx(path, limit))
    elif kind == 'word-legacy':
        base['warnings'].append('旧版 .doc 格式暂不支持内容提取，建议用 Word 另存为 .docx 后再查看。')
    elif kind == 'excel':
        if ext in ('xlsx', 'xlsm'):
            base.update(_extract_xlsx(path, limit))
        else:
            base['warnings'].append('旧版 .xls / .et 格式暂不支持内容提取。')
    elif kind == 'ppt':
        base.update(_extract_pptx(path, limit))
    elif kind == 'text':
        base.update(_extract_text(path, limit))
    elif kind == 'zip':
        base.update(_extract_zip(path, limit))
    elif kind == 'image':
        base['image_size'] = _image_size(path)
        base['warnings'].append('图片文件，请在预览中查看。')
    elif kind == 'archive':
        base['warnings'].append('RAR/7z 等压缩格式暂不支持列出内容。')
    elif kind == 'media':
        base['warnings'].append('音视频文件，请用本地播放器打开。')
    else:
        base['warnings'].append('二进制文件，仅提供基本信息。')

    if not base.get('title'):
        base['title'] = os.path.splitext(os.path.basename(path))[0]
    base['full_available'] = base.get('content_truncated', False) and not full
    return base
